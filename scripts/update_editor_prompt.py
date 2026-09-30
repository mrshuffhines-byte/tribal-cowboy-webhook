import copy
from docx import Document
from docx.text.paragraph import Paragraph

P = "/Users/staciehuffhines/Documents/Claude Yap Edits ready to edit/Yap Edits files/Yap_Edits_Editor_Prompt.docx"
d = Document(P)
paras = d.paragraphs
if any(p.text.startswith("10. VIDEO TYPES") for p in paras):
    raise SystemExit("section 10 already present; nothing to do")
heading_src = next(p for p in paras if p.text.startswith("9. TRIBAL COWBOY RULES"))
bullet_src = next(p for p in paras if p.text.startswith("- Background:"))
blank_src = next(p for p in paras if p.text == "")
body = d.element.body
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

def add(src, text):
    new = copy.deepcopy(src._p)
    body.append(new)
    par = Paragraph(new, src._parent)
    if par.runs:
        par.runs[0].text = text
        for r in par.runs[1:]:
            r.text = ""
    elif text:
        par.add_run(text)

add(blank_src, "")
add(heading_src, "10. VIDEO TYPES")
lines = [
    "- Before you edit anything, list the files you found with their durations, then ask me three things and wait for my answer: Which type: A talking, B voice over, C highlight, D trend? How long should it run: under 30 s, 30 to 60 s, or 60 to 90 s? Setting: outdoor, truck, or indoor? Do nothing until I reply.",
    "- Type A, talking reel: one clip of me talking to camera. Everything in sections 1 to 9 applies. Pick a Playbook format (A to D) and tell me why.",
    "- Type B, voice over reel: a subfolder holding 3 to 10 clips plus one voice memo (.m4a, .mp3, .caf, .wav, .aac, any capitalization). Transcribe the memo with word timestamps. Order the clips so each sentence lands on the footage it describes: a horse's name lands on that horse, a place name on that place. Cut each clip to the length of its sentence. If a clip is shorter than its sentence, slow push in 1.0 to 1.08x; never freeze a frame. Silence trimming applies to the voice, not the footage. One word captions sit in the bottom third, centered, since there is no chin to anchor to. Headlines and notes sit in the upper third, below the top 15 percent. Series tag on frame one. Follow ending. If no voice memo is in the folder, ask me for it; never narrate with a synthetic voice.",
    "- Type C, highlight cut: one clip longer than 60 seconds. Scan it for the 5 to 15 second stretches with the most motion, faces, or reaction sound. Return 3 to 5 separate vertical clips named highlight-01.mp4, highlight-02.mp4 and so on. Each starts with a text hook on frame one, keeps the natural sound, and ends with the Follow button. Write highlights.txt with one line per clip: file name, start and end time in the source, and a six word description.",
    "- Type D, trend remix: I give you a trend brief entry (the trend, the ranch version, the hook, the length, the beat structure) and the clip I filmed for it. Edit to the length and beats in the brief. Leave a 2 second lead at the start and export with no music; the sound gets added in Instagram or TikTok when I post.",
    "- Output for every type: a folder Edited/YYYY-MM-DD-slug/ holding render.mp4 (or the highlight files), transcript.txt, overlay-plan.md, and type.txt containing the single type letter. Tell me the folder name when you finish.",
]
for t in lines:
    add(bullet_src, t)
sect = body.find(W + "sectPr")
if sect is not None:
    body.remove(sect)
    body.append(sect)
d.save(P)
print("section 10 added")
